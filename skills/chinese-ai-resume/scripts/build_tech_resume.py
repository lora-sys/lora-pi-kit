#!/usr/bin/env python3
"""Build a one-page Chinese ATS-friendly AI-resume DOCX from JSON data."""

import argparse
import json
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = "173B70"
BLUE = "3568ED"
INK = "1B2A41"
MUTED = "617187"
PALE = "EDF3FF"
TABLE_TINT = "F7FAFF"


def set_font(run, size, bold=False, color=INK):
    run.font.name = "等线"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    props = run._element.get_or_add_rPr()
    fonts = props.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.append(fonts)
    fonts.set(qn("w:eastAsia"), "等线")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")


def space(paragraph, before=0, after=0, line=1.0, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.keep_with_next = keep


def clear_borders(table):
    props = table._tbl.tblPr
    borders = props.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        props.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "nil")
        borders.append(node)


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    props.append(node)


def cell_margin(cell, top=16, left=36, bottom=16, right=36):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side, value in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        item = margins.find(qn(f"w:{side}"))
        if item is None:
            item = OxmlElement(f"w:{side}")
            margins.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def link(paragraph, label, url, size=6.8, color=BLUE):
    relation_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    props.append(color_node)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "等线")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    props.append(fonts)
    size_node = OxmlElement("w:sz")
    size_node.set(qn("w:val"), str(int(size * 2)))
    props.append(size_node)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def section(doc, name):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    clear_borders(table)
    table.columns[0].width = Cm(0.16)
    table.columns[1].width = Cm(18.25)
    stripe, title = table.rows[0].cells
    shade(stripe, BLUE)
    shade(title, PALE)
    cell_margin(title, 18, 62, 18, 32)
    paragraph = title.paragraphs[0]
    space(paragraph, keep=True)
    run = paragraph.add_run(name)
    set_font(run, 8.9, True, NAVY)
    gap = doc.add_paragraph()
    space(gap, after=0.14, line=0.12, keep=True)


def bullet(doc, text, number=None, size=6.8):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.39)
    paragraph.paragraph_format.first_line_indent = Cm(-0.39)
    space(paragraph, after=0.01, line=0.98, keep=True)
    marker = paragraph.add_run(f"{number}. " if number else "• ")
    set_font(marker, size, True, BLUE)
    run = paragraph.add_run(text)
    set_font(run, size, False, INK)


def subsection(doc, name):
    paragraph = doc.add_paragraph()
    space(paragraph, line=1, keep=True)
    run = paragraph.add_run(name)
    set_font(run, 7.1, True, NAVY)


def title_row(doc, title, date, title_size=7.9):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    clear_borders(table)
    table.columns[0].width = Cm(14.7)
    table.columns[1].width = Cm(3.7)
    left, right = table.rows[0].cells
    p = left.paragraphs[0]
    space(p, line=1, keep=True)
    run = p.add_run(title)
    set_font(run, title_size, True, INK)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    space(p, line=1, keep=True)
    run = p.add_run(date)
    set_font(run, 6.8, False, MUTED)


def experience(doc, item):
    title_row(doc, item["title"], item.get("date", ""))
    if item.get("subject"):
        p = doc.add_paragraph()
        space(p, before=0.01, after=0.01, line=0.98, keep=True)
        run = p.add_run("负责项目：")
        set_font(run, 6.8, True, BLUE)
        run = p.add_run(item["subject"])
        set_font(run, 6.8, False, INK)
    for text in item.get("bullets", []):
        bullet(doc, text, size=6.75)


def project(doc, item):
    title_row(doc, item["title"], item.get("date", ""), title_size=8.85)
    p = doc.add_paragraph()
    space(p, after=0.015, line=0.96, keep=True)
    run = p.add_run(item.get("link_label", "GitHub 开源") + "：")
    set_font(run, 6.8, True, BLUE)
    link(p, item.get("url_label", item["url"].replace("https://", "")), item["url"], 6.8, NAVY)
    p = doc.add_paragraph()
    space(p, after=0.015, line=0.98, keep=True)
    run = p.add_run("项目介绍：")
    set_font(run, 6.9, True, BLUE)
    run = p.add_run(item["intro"])
    set_font(run, 6.9, False, INK)
    p = doc.add_paragraph()
    space(p, after=0.005, line=0.96, keep=True)
    run = p.add_run("主要工作：")
    set_font(run, 6.9, True, BLUE)
    for index, text in enumerate(item.get("bullets", []), start=1):
        bullet(doc, text, index, 6.7)
    gap = doc.add_paragraph()
    space(gap, after=0.025, line=0.08, keep=True)


def work_item(table, item, tinted=False):
    cells = table.add_row().cells
    for cell in cells:
        cell_margin(cell, 8, 26, 8, 26)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if tinted:
            shade(cell, TABLE_TINT)
    p = cells[0].paragraphs[0]
    space(p, line=1)
    run = p.add_run(item["name"])
    set_font(run, 6.8, True, NAVY)
    p = cells[1].paragraphs[0]
    space(p, line=1)
    run = p.add_run(item["description"])
    set_font(run, 6.65, False, INK)
    p = cells[2].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    space(p, line=1)
    link(p, item.get("link_label", "访问"), item["url"], 6.6, BLUE)


def build(data, output):
    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin = page.bottom_margin = Cm(0.36)
    page.left_margin = page.right_margin = Cm(1.14)
    page.footer_distance = Cm(0.18)
    normal = doc.styles["Normal"]
    normal.font.name = "等线"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(7)
    doc.core_properties.title = f'{data["name"]}｜{data["target"]}'
    doc.core_properties.author = data["name"]

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, after=0.12, line=1, keep=True)
    run = p.add_run(data["name"])
    set_font(run, 19, True, NAVY)
    run = p.add_run("  ｜  " + data["target"])
    set_font(run, 8.2, True, BLUE)
    for line in data.get("contact_lines", []):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        space(p, after=0.01, line=1, keep=True)
        run = p.add_run(line)
        set_font(run, 6.8, False, MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    space(p, after=0.16, line=1, keep=True)
    for index, contact in enumerate(data.get("profile_links", [])):
        if index:
            run = p.add_run("  ｜  ")
            set_font(run, 6.8, False, MUTED)
        run = p.add_run(contact["label"] + "：")
        set_font(run, 6.8, False, MUTED)
        link(p, contact["text"], contact["url"], 6.8, NAVY)

    education = data.get("education")
    if education:
        section(doc, "教育经历")
        p = doc.add_paragraph()
        space(p, after=0.01, line=0.98, keep=True)
        run = p.add_run(education["school"])
        set_font(run, 7.55, True, INK)
        run = p.add_run("  " + education.get("detail", ""))
        set_font(run, 6.8, False, MUTED)
        if education.get("honors"):
            p = doc.add_paragraph()
            space(p, after=0.12, line=0.98, keep=True)
            run = p.add_run(education["honors"])
            set_font(run, 6.65, False, MUTED)

    if data.get("skills"):
        section(doc, "专业技能")
        for group in data["skills"]:
            subsection(doc, group["name"])
            for text in group.get("bullets", []):
                bullet(doc, text, size=6.7)

    if data.get("experience"):
        section(doc, "工作经历")
        for item in data["experience"]:
            experience(doc, item)

    if data.get("projects"):
        section(doc, "项目经历")
        for item in data["projects"]:
            project(doc, item)

    if data.get("works"):
        section(doc, "其他作品")
        table = doc.add_table(rows=0, cols=3)
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        clear_borders(table)
        table.columns[0].width, table.columns[1].width, table.columns[2].width = Cm(3.15), Cm(11.55), Cm(3.55)
        for index, item in enumerate(data["works"]):
            work_item(table, item, tinted=(index == 0))

    if data.get("advantages"):
        section(doc, "个人优势")
        for text in data["advantages"]:
            bullet(doc, text, size=6.75)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def convert_to_pdf(docx_path, pdf_output):
    office = shutil.which("libreoffice") or shutil.which("soffice")
    if not office:
        raise RuntimeError("未找到 LibreOffice 或 soffice，无法转换 PDF。")
    temp_dir = pdf_output.parent / ".resume-pdf-temp"
    temp_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", str(temp_dir), str(docx_path)], check=True)
    converted = temp_dir / f"{docx_path.stem}.pdf"
    if not converted.exists():
        raise RuntimeError("LibreOffice 未生成 PDF。")
    shutil.move(str(converted), str(pdf_output))
    temp_dir.rmdir()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, help="resume JSON file")
    parser.add_argument("--docx", required=True, help="output DOCX path")
    parser.add_argument("--pdf", help="optional output PDF path")
    args = parser.parse_args()
    with Path(args.input).open("r", encoding="utf-8") as file:
        data = json.load(file)
    docx_path = Path(args.docx).resolve()
    build(data, docx_path)
    print(f"DOCX: {docx_path}")
    if args.pdf:
        pdf_path = Path(args.pdf).resolve()
        convert_to_pdf(docx_path, pdf_path)
        print(f"PDF: {pdf_path}")


if __name__ == "__main__":
    main()
